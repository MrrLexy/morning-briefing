#!/usr/bin/env python3
"""
Generates bbg_template.xlsx — the Excel workbook you refresh each morning.

The workbook is a grid of =BDP() formulas, one row per security, one column per
Bloomberg field. Open it on a machine with the Bloomberg Excel add-in and the
formulas resolve to live values.

Run this only when the security list or field set changes:

    python make_bbg_template.py

Keeping the template generated (rather than hand-maintained) means the tickers in
it can't silently drift from SECURITIES in fetch_bloomberg.py and watchlist.json,
which are the source of truth.
"""

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from fetch_bloomberg import FIELDS, SECURITIES, WATCHLIST_FILE

OUT_PATH = Path(__file__).parent / "bbg_template.xlsx"

# Internal shorthand -> real Bloomberg syntax (what BDP needs).
# Rule is suffix-based, but USGGT10Y is filed under :GOV internally while
# Bloomberg exposes it as an Index, so it needs an explicit entry.
SUFFIX_TO_YELLOW_KEY = {
    "GOV": "Govt",
    "IND": "Index",
    "CUR": "Curncy",
    "COM": "Comdty",
    "US": "US Equity",
}
EXPLICIT = {
    "USGGT10Y:GOV": "USGGT10Y Index",
}

# Which fields actually apply to which kind of security. Requesting a field a
# security doesn't have just yields #N/A — harmless, the importers treat it as
# blank — but a template full of #N/A is alarming to read, so scope it.
BENCHMARK_FIELDS = ["PX_LAST", "CHG_NET_1D", "YLD_YTM_BID", "OAS_SPREAD_ASK",
                    "DUR_ADJ_MID", "RETURN_YTD", "RETURN_QTD"]
TREASURY_FIELDS = ["PX_LAST", "CHG_NET_1D", "YLD_YTM_BID"]
LEVEL_FIELDS = ["PX_LAST", "CHG_NET_1D"]
CROSS_ASSET_FIELDS = ["PX_LAST", "CHG_PCT_1D"]
ETF_FIELDS = ["PX_LAST", "CHG_PCT_1D", "YLD_YTM_BID", "OAS_SPREAD_ASK", "RETURN_YTD"]

BENCHMARKS = {"LBUSTRUU:IND", "LUACTRUU:IND", "JBCDCOMP:IND", "BAMLH0A1HYBB:IND"}
CROSS_ASSET = {"SPX:IND", "DXY:CUR", "CL1:COM"}


def to_bloomberg_syntax(shorthand: str) -> str:
    if shorthand in EXPLICIT:
        return EXPLICIT[shorthand]
    base, _, suffix = shorthand.partition(":")
    return f"{base} {SUFFIX_TO_YELLOW_KEY.get(suffix, 'Index')}"


def fields_for(shorthand: str) -> list:
    if shorthand in BENCHMARKS:
        return BENCHMARK_FIELDS
    if shorthand in CROSS_ASSET:
        return CROSS_ASSET_FIELDS
    if shorthand.endswith(":GOV"):
        return TREASURY_FIELDS
    if shorthand.endswith(":US"):
        return ETF_FIELDS
    return LEVEL_FIELDS


def load_watchlist_tickers() -> list:
    if not WATCHLIST_FILE.exists():
        return []
    cfg = json.loads(WATCHLIST_FILE.read_text())
    return [t["ticker"] for t in cfg.get("tickers", []) if t.get("ticker")]


def build() -> None:
    watchlist = load_watchlist_tickers()
    tickers = SECURITIES + [t for t in watchlist if t not in SECURITIES]

    wb = Workbook()
    ws = wb.active
    ws.title = "Data"

    header_fill = PatternFill("solid", fgColor="003057")
    header_font = Font(color="FFFFFF", bold=True, size=10)
    thin = Side(style="thin", color="D6DAE4")
    border = Border(bottom=thin, right=thin)

    headers = ["Ticker"] + FIELDS
    for col, name in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    for row_i, shorthand in enumerate(tickers, start=2):
        bbg = to_bloomberg_syntax(shorthand)
        applicable = fields_for(shorthand)

        c = ws.cell(row=row_i, column=1, value=bbg)
        c.font = Font(size=10)
        c.border = border

        for col_i, field in enumerate(FIELDS, start=2):
            cell = ws.cell(row=row_i, column=col_i)
            if field in applicable:
                cell.value = f'=BDP("{bbg}","{field}")'
            cell.font = Font(size=10)
            cell.border = border
            cell.alignment = Alignment(horizontal="right")

    ws.column_dimensions["A"].width = 22
    for col_i in range(2, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col_i)].width = 15
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "B2"

    # Instructions on a second sheet so they can't be mistaken for data rows.
    notes = wb.create_sheet("How to use")
    lines = [
        ("Morning Briefing — Bloomberg data template", True),
        ("", False),
        ("Each morning:", True),
        ("1. Open this file on a machine with the Bloomberg Excel add-in.", False),
        ("2. Wait for the =BDP() formulas on the Data sheet to resolve to numbers.", False),
        ("   #N/A in a cell is fine — that field doesn't apply to that security.", False),
        ("3. File > Save As > CSV (Comma delimited). Saving as CSV converts the", False),
        ("   formulas to plain values automatically — that IS the 'save as values' step.", False),
        ("4. Feed the CSV into the briefing, either:", False),
        ("      python import_bbg_paste.py yourfile.csv", False),
        ("   or the browser converter (no Python needed):", False),
        ("      https://mrrlexy.github.io/morning-briefing/import/", False),
        ("", False),
        ("Keep this .xlsx — it is the reusable template. Save the CSV under a new", False),
        ("name each day so you never overwrite the formulas.", False),
        ("", False),
        ("Do not commit the CSV or data.json to git. Bloomberg values stay on this", True),
        ("machine — see DATA_POLICY.md.", True),
        ("", False),
        ("To change which securities or fields appear here, edit SECURITIES in", False),
        ("fetch_bloomberg.py or watchlist.json, then re-run make_bbg_template.py.", False),
    ]
    for i, (text, bold) in enumerate(lines, start=1):
        cell = notes.cell(row=i, column=1, value=text)
        cell.font = Font(bold=bold, size=11)
    notes.column_dimensions["A"].width = 95

    wb.save(OUT_PATH)
    print(f"  OK: Wrote {OUT_PATH}")
    print(f"  OK: {len(tickers)} securities x {len(FIELDS)} fields")


if __name__ == "__main__":
    build()
